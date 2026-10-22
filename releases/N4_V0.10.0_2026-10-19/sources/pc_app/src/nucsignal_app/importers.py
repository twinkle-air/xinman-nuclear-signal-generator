from __future__ import annotations

import csv
import hashlib
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from .models import SignalSource


BUILTIN_NUCLIDES = (
    "Am-241", "Cs-137", "Co-60", "K-40",
    "Na-22", "Ba-133", "Eu-152", "I-131", "Tc-99m",
)


def _hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _read_rows(path: Path) -> list[list[object]]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        return rows

    text = path.read_text(encoding="utf-8-sig")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t ")
        return [row for row in csv.reader(text.splitlines(), dialect) if row]
    except csv.Error:
        return [line.replace(",", " ").split() for line in text.splitlines() if line.strip()]


def _numeric_rows(rows: list[list[object]]) -> np.ndarray:
    parsed: list[tuple[float, float]] = []
    for row_number, row in enumerate(rows, start=1):
        if len(row) < 2:
            continue
        try:
            parsed.append((float(row[0]), float(row[1])))
        except (TypeError, ValueError):
            if row_number == 1:
                continue
            raise ValueError(f"第 {row_number} 行不是有效的两列数值") from None
    if not parsed:
        raise ValueError("文件中没有可用的两列数值")
    return np.asarray(parsed, dtype=np.float64)


def load_spectrum(path: str | Path) -> SignalSource:
    source_path = Path(path)
    data = _numeric_rows(_read_rows(source_path))
    if data.shape[0] not in {256, 1024}:
        raise ValueError(f"能谱必须是 256 或 1024 道，当前为 {data.shape[0]} 道")
    if np.any(data[:, 1] < 0) or not np.isfinite(data).all():
        raise ValueError("能谱计数必须是有限的非负数")
    if data[:, 1].sum() <= 0:
        raise ValueError("能谱计数总和必须大于 0")
    if np.unique(data[:, 0]).size != data.shape[0]:
        raise ValueError("能谱存在重复道号/能量")
    probability = data[:, 1] / data[:, 1].sum()
    return SignalSource(
        name=source_path.stem,
        kind="spectrum",
        x=data[:, 0],
        y=probability,
        source_path=source_path,
        sha256=_hash_file(source_path),
        full_scale_x=float(np.max(data[:, 0])),
    )


def load_waveform(path: str | Path, sample_rate_hz: float = 50e6) -> SignalSource:
    source_path = Path(path)
    data = _numeric_rows(_read_rows(source_path))
    if data.shape[0] < 2:
        raise ValueError("波形至少需要两个采样点")
    if not np.isfinite(data).all() or np.any(np.diff(data[:, 0]) <= 0):
        raise ValueError("时间列必须有限且严格递增")
    # 输入时间单位统一定义为微秒。
    target_time = np.arange(data[0, 0], data[-1, 0] + 1e-12, 1e6 / sample_rate_hz)
    if target_time.size > 5000:
        raise ValueError("重采样后超过 5000 点（100 μs），请截断或分段导入")
    normalized = np.interp(target_time, data[:, 0], data[:, 1])
    normalized = np.clip(normalized, -1.0, 1.0)
    quantized = np.round(normalized * 8191) / 8191
    return SignalSource(
        name=source_path.stem,
        kind="waveform",
        x=target_time,
        y=quantized,
        source_path=source_path,
        sha256=_hash_file(source_path),
    )


def builtin_spectrum(name: str, channels: int = 1024) -> SignalSource:
    energies = np.linspace(0, 2000, channels, endpoint=False)

    def gaussian(center: float, fwhm: float, weight: float = 1.0) -> np.ndarray:
        sigma = fwhm / 2.35482
        return weight * np.exp(-0.5 * ((energies - center) / sigma) ** 2)

    if name == "Am-241":
        counts = gaussian(59.54, 12.0) + 0.04 * np.exp(-energies / 170)
        reference_peaks = (59.54,)
    elif name == "Cs-137":
        edge = np.where(energies < 478, 0.10 * np.sqrt(np.maximum(478 - energies, 0) / 478), 0)
        counts = gaussian(661.657, 52.0) + edge
        reference_peaks = (661.657,)
    elif name == "Co-60":
        counts = gaussian(1173.228, 82.0) + gaussian(1332.492, 88.0, 0.96)
        reference_peaks = (1173.228, 1332.492)
    elif name == "K-40":
        counts = gaussian(1460.8, 95.0)
        reference_peaks = (1460.8,)
    elif name == "Na-22":
        counts = gaussian(511.0, 45.0, 1.8) + gaussian(1274.537, 86.0)
        reference_peaks = (511.0, 1274.537)
    elif name == "Ba-133":
        counts = (
            gaussian(80.997, 13.0, 0.34) + gaussian(276.398, 29.0, 0.072)
            + gaussian(302.851, 31.0, 0.183) + gaussian(356.013, 34.0, 0.62)
            + gaussian(383.849, 36.0, 0.089)
        )
        reference_peaks = (80.997, 302.851, 356.013)
    elif name == "Eu-152":
        counts = (
            gaussian(121.782, 15.0, 0.286) + gaussian(244.697, 27.0, 0.075)
            + gaussian(344.279, 34.0, 0.266) + gaussian(778.905, 59.0, 0.129)
            + gaussian(964.079, 67.0, 0.146) + gaussian(1112.074, 76.0, 0.135)
            + gaussian(1408.013, 92.0, 0.208)
        )
        reference_peaks = (121.782, 344.279, 1408.013)
    elif name == "I-131":
        counts = (
            gaussian(284.305, 30.0, 0.061) + gaussian(364.489, 36.0, 0.812)
            + gaussian(636.989, 51.0, 0.072) + gaussian(722.911, 56.0, 0.018)
        )
        reference_peaks = (364.489,)
    elif name == "Tc-99m":
        counts = gaussian(140.511, 17.0)
        reference_peaks = (140.511,)
    else:
        raise ValueError(f"未知核素模板：{name}")
    counts = np.maximum(counts, 0)
    return SignalSource(
        name=name,
        kind="spectrum",
        x=energies,
        y=counts / counts.sum(),
        full_scale_x=2000.0,
        reference_peaks=reference_peaks,
    )

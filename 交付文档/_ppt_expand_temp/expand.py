from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(r"C:\Users\Twinkle\Desktop\新漫杯")
SOURCE = ROOT / "交付文档" / "智能仿核信号发生器_新漫杯决赛答辩.pptx"
OUTPUT = ROOT / "交付文档" / "_ppt_expand_temp" / "candidate.pptx"
prs = Presentation(str(SOURCE))
assert len(prs.slides) == 9
NAVY = RGBColor(24, 56, 89)
BLUE = RGBColor(30, 106, 160)
TEAL = RGBColor(14, 143, 143)
GREEN = RGBColor(43, 150, 112)
PALE = RGBColor(239, 246, 250)
PALE_GREEN = RGBColor(237, 248, 244)
INK = RGBColor(34, 50, 65)
MID = RGBColor(77, 94, 109)
LINE = RGBColor(206, 222, 230)
WHITE = RGBColor(255, 255, 255)
AMBER = RGBColor(166, 90, 12)
PALE_AMBER = RGBColor(255, 249, 237)


def txt(s, x, y, w, h, value, size=15, bold=False, color=INK,
        align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.055)
    tf.margin_top = tf.margin_bottom = Inches(.015)
    tf.vertical_anchor = valign
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_before = p.space_after = Pt(0)
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
    return sh


def box(s, x, y, w, h, fill=WHITE, stroke=LINE, round=True):
    k = MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE
    sh = s.shapes.add_shape(k, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = stroke
    sh.line.width = Pt(1)
    return sh


def ln(s, x1, y1, x2, y2, color=TEAL, width=2):
    sh = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = color
    sh.line.width = Pt(width)
    return sh


def new(title, section):
    s = prs.slides.add_slide(prs.slide_layouts[2])
    s.shapes[0].text = title
    txt(s, .56, 1.33, 12.15, .35, section, 16, True, BLUE)
    return s


def foot(s, value):
    txt(s, .52, 7.01, 12.26, .28, value, 9, False, MID)


def card(s, x, y, w, h, title, body, color=BLUE, fill=WHITE, tsize=15, bsize=12):
    box(s, x, y, w, h, fill)
    txt(s, x+.18, y+.17, w-.36, .40, title, tsize, True, color)
    txt(s, x+.18, y+.68, w-.36, h-.83, body, bsize, False, INK,
        valign=MSO_ANCHOR.TOP)


# Inserted before the architecture page: requirement decomposition and current scope.
s4 = new("赛题指标与当前作品边界", "赛题要求拆解：先达成可演示的基本链路，再扩展全指标覆盖")
card(s4, .56, 1.90, 5.96, 3.97, "赛题关注的能力",
     "· 可配置核脉冲与模拟电压输出\n· 幅度、半高宽、周期及噪声参数调节\n· 上位机交互与串口控制\n· 进阶：自定义能谱、随机到达、多通道等\n· 通过测试数据说明精度和稳定性",
     BLUE, PALE, 17, 14)
card(s4, 6.76, 1.90, 6.00, 3.97, "当前样机的实际完成范围",
     "· PC—STM32—FPGA—AD9708 单通道链路\n· 单指数脉冲、固定周期及随机事件定性演示\n· 代表幅度/宽度/周期点有示波器记录\n· 20%FS 高斯设定已发现削顶\n· 双指数、双通道与全量程校准尚未完成",
     GREEN, PALE_GREEN, 17, 14)
box(s4, .56, 6.15, 12.20, .50, PALE_AMBER, PALE_AMBER)
txt(s4, .74, 6.22, 11.84, .33,
    "答辩口径：能演示且有局部实测证据 ≠ 赛题全部指标通过。", 13, True, AMBER)
foot(s4, "赛题目标与样机阶段分开陈述，避免将后续升级功能计为已完成成果。")


# Extra hardware page.
s6 = new("硬件模块与接口划分", "第一阶段采用现成模块联调，接口、供电和模拟输出分层")
modules = [
    (0.56, "上位机 + USB 转串口", "CH340 / USART1\n参数帧、握手与状态", BLUE),
    (3.66, "STM32F405 主控", "合法性校验\nSPI 控制与停止路径", TEAL),
    (6.76, "G1 FPGA 逻辑", "50 MHz 时钟\n事件/波形/DAC 时序", BLUE),
    (9.86, "高速 ADDA 模块", "AD9708 8 bit DAC\nBNC 输出到示波器", TEAL),
]
for x, a, b, c in modules:
    card(s6, x, 2.05, 2.90, 1.87, a, b, c, PALE, 14, 12)
for x in (3.46, 6.56, 9.66):
    ln(s6, x, 2.96, x+.17, 2.96)
box(s6, .56, 4.30, 12.20, 2.10, WHITE)
txt(s6, .81, 4.52, 11.72, .36, "关键接口与设计约束", 16, True, NAVY)
txt(s6, .82, 5.04, 11.57, 1.13,
    "PC ↔ STM32：串口帧携带配置与状态；STM32 ↔ FPGA：SPI 下发参数并进行提交/启动；FPGA → DAC：并行采样码与更新时钟。\n模拟输出需关注电源噪声、数字地回流、线缆与示波器输入设置；现阶段没有完成模拟反馈闭环校准。",
    13, False, INK)
foot(s6, "现有实物为 ATK-HS-ADDA V3.3，商家确认其对应 V3.2 原理图；AD9767 双通道留作第二阶段。")


# Extra communication state machine page.
s8 = new("通讯协作与安全状态", "将配置、提交、启动、读状态、停止分成可观察的步骤")
steps = [
    (.56, "HELLO", "联机/版本"), (2.58, "SET", "校验参数"),
    (4.60, "COMMIT", "装载配置"), (6.62, "START", "开始输出"),
    (8.64, "STATUS", "监测状态"), (10.66, "STOP", "回到空闲"),
]
for x, a, b in steps:
    box(s8, x, 2.25, 1.73, 1.35, PALE)
    txt(s8, x+.08, 2.44, 1.57, .36, a, 15, True, NAVY, PP_ALIGN.CENTER)
    txt(s8, x+.08, 2.95, 1.57, .31, b, 11, False, MID, PP_ALIGN.CENTER)
for x in (2.29, 4.31, 6.33, 8.35, 10.37):
    ln(s8, x, 2.93, x+.25, 2.93)
card(s8, .56, 4.03, 5.95, 2.38, "正常控制链",
     "PC 请求 → STM32 校验 → FPGA 接收活动参数 → DAC 输出。每一步读取应答/状态，参数异常拒绝提交；避免只凭界面显示推断硬件已更新。",
     BLUE, WHITE, 16, 13)
card(s8, 6.78, 4.03, 5.98, 2.38, "停止与失联处理",
     "现有台架有 PA0 停止、关闭上位机后停止及重连不自启记录。N3 的 30 分钟人工运行记录不能直接等同 N4 当前版本长时可靠性。",
     GREEN, PALE_GREEN, 16, 13)
foot(s8, "真实 A/B/PG 安全输入与全部异常场景未完成独立验收；展示时不称为完整安全联锁。")


# Extra waveform-parameter page.
s9 = new("参数到波形的生成路径", "把用户参数映射为事件间隔、脉冲形状与 DAC 码")
card(s9, .56, 1.92, 3.92, 3.34, "① 事件何时到达",
     "固定周期：按设定间隔触发。\n随机到达：离散近似生成不等间隔事件；到达率分布尚待定量统计。",
     BLUE, PALE, 16, 14)
card(s9, 4.70, 1.92, 3.92, 3.34, "② 单次脉冲多高、多宽",
     "幅度来自常量或 256 点谱分位表；单指数离散衰减决定形状。τ 与半高宽不是同一个量，需要由实测 FWHM 核对。",
     TEAL, PALE_GREEN, 16, 14)
card(s9, 8.84, 1.92, 3.92, 3.34, "③ 如何成为模拟信号",
     "背景与高斯近似噪声在数字域叠加，量化为 8 bit DAC 码，以约 1 μs 步进更新；超出码域会削顶。",
     BLUE, PALE, 16, 14)
box(s9, .56, 5.62, 12.20, .89, WHITE)
txt(s9, .83, 5.79, 11.65, .49,
    "波形设计变量：事件间隔 T、幅度 A、衰减 τ、基线 B、噪声强度 σ；测试量是实际 Vpp/FWHM/周期/RMS。",
    14, True, NAVY)
foot(s9, "“设置值”和“输出实测值”在演示与报告中分别标注，不能互相替代。")


# Extra testing methodology page.
s10 = new("测试方法与证据链", "同一测试点保留配置、示波器记录与分析结果")
cards = [
    (.56, "配置记录", "记录 PC/STM32/FPGA 版本、幅度/宽度/周期/噪声设定及活动参数截图。"),
    (4.70, "波形采集", "Hantek DSO2D15 示波器，BNC 直连 CH2、X1、DC；截图与原始 CSV 同时归档。"),
    (8.84, "离线核验", "从 CSV 计算局部基线、峰值、半高宽、周期与 RMS；结果与截图交叉复核。"),
]
for x, a, b in cards:
    card(s10, x, 2.12, 3.92, 2.55, a, b, BLUE, PALE, 16, 13)
box(s10, .56, 5.02, 12.20, 1.33, PALE_AMBER, PALE_AMBER)
txt(s10, .80, 5.22, 11.72, .84,
    "证据边界：仪器已做自校准，但无独立有效校准参考；短时单点工程读数不能推出全量程、可溯源绝对准确度。", 14, True, AMBER)
foot(s10, "2026-10-22 五组记录附原始 CSV、截图与 SHA256；未从单次截图推断长时间稳定性。")


# Extra detailed metrics page.
s12 = new("关键性能分项核验", "代表点读数与赛题指标对照：只对已测点给结论")
box(s12, .56, 1.88, 12.20, 4.92, WHITE)
heads = [(0.87, "项目"), (4.19, "设定/目标"), (6.82, "实测"), (9.46, "说明")]
for x, h in heads:
    txt(s12, x, 2.12, 2.7, .38, h, 13, True, NAVY)
ln(s12, .82, 2.57, 12.46, 2.57, LINE, 1)
rows = [
    ("幅度代表点", "3.3 V", "3.260 V", "−1.21%，单点"),
    ("最小脉宽", "1 μs FWHM", "1.0285 μs", "+2.85%，单脉冲"),
    ("最大脉宽", "100 μs FWHM", "98.00 μs", "−2.00%，单点"),
    ("最短周期", "10 μs", "10.000 μs", "100 kHz，短记录"),
    ("最长周期", "10 ms", "10.000 ms", "单点"),
]
for i, row in enumerate(rows):
    yy = 2.71 + i*.64
    if i%2 == 0:
        box(s12, .82, yy-.03, 11.67, .55, PALE, PALE, False)
    for (x,_), val, w in zip(heads, row, (3.02, 2.36, 2.30, 2.83)):
        txt(s12, x, yy, w, .43, val, 12.5, i==0 and x==6.82, INK)
txt(s12, .84, 6.14, 11.50, .39,
    "单点数值可支持阶段性演示，但尚未覆盖每 0.1 V、每 1 μs、全部周期档位及参数组合。", 12.5, True, AMBER)
foot(s12, "最短周期 10 μs 来自 6 个完整内部间隔；不将该短记录外推为持续 100 kHz 长时稳定性。")


# Extra noise and qualitative results page.
s13 = new("噪声、能谱与连续运行记录", "展示效果、统计读数与版本差异分别标注")
card(s13, .56, 1.96, 3.92, 3.89, "高斯噪声：有量化也有边界",
     "5%FS 三帧 AC RMS：312.03 / 302.21 / 314.90 mV，均值 309.72 mV。\n\n20%FS 设置：AC RMS 1.232 V，设备累计削顶 90,879 点；不能称无削顶达标。",
     BLUE, PALE, 16, 13)
card(s13, 4.70, 1.96, 3.92, 3.89, "能谱：粗分布演示",
     "Cs-137 内置量化模板下，10 帧共 615 个完整周期，脉高粗分布与本地模板接近。\n\n这是仿真波形表现，不是实测核素源谱的峰位或能量分辨率验证。",
     TEAL, PALE_GREEN, 16, 13)
card(s13, 8.84, 1.96, 3.92, 3.89, "连续运行：保留版本边界",
     "2026-10-19 用户人工记录 N3 台架运行 30 分钟无通信/波形异常，停止与重连流程正常。\n\n不可直接作为当前 N4 版本的长时间稳定性证明。",
     BLUE, PALE, 16, 13)
box(s13, .56, 6.08, 12.20, .47, PALE_AMBER, PALE_AMBER)
txt(s13, .77, 6.14, 11.77, .33,
    "仍待做：噪声逐档/带宽、随机到达分布、精密能谱和当前版本长时漂移验证。", 12, True, AMBER)
foot(s13, "不同日期/版本的证据不可拼接为同一轮整机全指标验收。")


# Extra innovation/roadmap page.
s16 = new("后续迭代与答辩重点", "从单通道基本链路走向双通道进阶功能，按证据逐步推进")
stages = [
    (.56, "阶段一：当前样机", "单通道 / 单指数；PC—STM32—FPGA—AD9708 链路；代表点示波器核验。", GREEN),
    (4.70, "阶段二：功能扩展", "引入 AD9767 双通道并验证接口兼容；实现双指数、双通道协作与更丰富谱形。", BLUE),
    (8.84, "阶段三：全指标验收", "逐档校准幅度/宽度/周期/噪声，补随机统计、长时漂移及真实安全输入。", TEAL),
]
for x, title, body, color in stages:
    card(s16, x, 2.02, 3.92, 3.43, title, body, color, PALE, 16, 14)
for x in (4.49, 8.63):
    ln(s16, x, 3.73, x+.18, 3.73, TEAL, 2)
box(s16, .56, 5.78, 12.20, .67, PALE_GREEN, PALE_GREEN)
txt(s16, .76, 5.91, 11.78, .39,
    "答辩强调：架构分工清晰、原始证据可复核；对未覆盖项和发现的削顶问题如实说明。", 13, True, NAVY)
foot(s16, "升级功能是路线图，不在当前样机成果或成本合计中计入。")


# Reorder nine existing slides and eight new slides to exactly 17 pages.
new_order = [
    prs.slides[0], prs.slides[1], prs.slides[2], s4,
    prs.slides[3], s6, prs.slides[4], s8, s9, s10,
    prs.slides[5], s12, s13, prs.slides[6], prs.slides[7], s16,
    prs.slides[8],
]
ids = prs.slides._sldIdLst
id_by_slide_id = {int(e.id): e for e in list(ids)}
for e in list(ids):
    ids.remove(e)
for slide in new_order:
    ids.append(id_by_slide_id[slide.slide_id])
assert len(prs.slides) == 17
prs.save(str(OUTPUT))
print(OUTPUT)
